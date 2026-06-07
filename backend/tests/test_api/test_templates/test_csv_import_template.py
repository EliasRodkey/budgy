#!python3
"""
Tests for the GET /templates/csv-import endpoint.

The endpoint returns an in-memory .xlsx workbook (no DB access), so these tests
call the handler directly, parse the returned bytes with openpyxl, and
cross-validate the structure against the Python source-of-truth category enums.
"""
from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.api.templates.templates_router import (
    TEMPLATE_HEADERS,
    download_csv_import_template,
)
from backend.utils.analysis_utils import DetailedCategories, PrimaryCategories


@pytest.fixture
def workbook():
    response = download_csv_import_template()
    return load_workbook(BytesIO(response.body))


class TestDownloadCsvImportTemplate:
    def test_returns_xlsx_media_type_and_attachment_header(self):
        response = download_csv_import_template()
        assert response.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert response.headers["content-disposition"] == 'attachment; filename="budgy_import_template.xlsx"'

    def test_main_sheet_header_row_matches_template_headers(self, workbook):
        ws = workbook["Transactions"]
        assert [cell.value for cell in ws[1]] == TEMPLATE_HEADERS

    def test_main_sheet_has_example_rows(self, workbook):
        ws = workbook["Transactions"]
        assert ws.max_row > 1

    def test_categories_sheet_is_hidden(self, workbook):
        assert workbook["Categories"].sheet_state == "hidden"

    def test_categories_sheet_lists_all_primary_categories(self, workbook):
        ws = workbook["Categories"]
        values = [cell[0].value for cell in ws.iter_rows(min_row=2, min_col=1, max_col=1) if cell[0].value]
        assert values == [c.value for c in PrimaryCategories]

    def test_categories_sheet_lists_all_detailed_categories(self, workbook):
        ws = workbook["Categories"]
        values = [cell[0].value for cell in ws.iter_rows(min_row=2, min_col=2, max_col=2) if cell[0].value]
        assert values == [c.value for c in DetailedCategories]

    def test_primary_category_column_has_dropdown_validation(self, workbook):
        ws = workbook["Transactions"]
        validations = ws.data_validations.dataValidation
        primary_col_validations = [dv for dv in validations if "A2" in str(dv.sqref)]
        assert len(primary_col_validations) == 1
        assert primary_col_validations[0].type == "list"
        assert primary_col_validations[0].formula1 == "=PrimaryCategoryList"

    def test_detailed_category_column_has_dropdown_validation(self, workbook):
        ws = workbook["Transactions"]
        validations = ws.data_validations.dataValidation
        detailed_col_validations = [dv for dv in validations if "B2" in str(dv.sqref)]
        assert len(detailed_col_validations) == 1
        assert detailed_col_validations[0].type == "list"
        assert detailed_col_validations[0].formula1 == "=DetailedCategoryList"
