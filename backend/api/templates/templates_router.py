#!python3
"""
backend.api.templates.templates_router

Serves a downloadable Excel (.xlsx) template for CSV/Excel transaction imports,
pre-populated with the expected column headers, example rows, and dropdown data
validation for the Primary/Detailed Category columns sourced from the
authoritative category enums in analysis_utils.py.
"""
from io import BytesIO

from fastapi import APIRouter, Response
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from backend.utils.analysis_utils import DetailedCategories, PrimaryCategories
from backend.utils.api_utils import RouterPrefixes

router = APIRouter(prefix=RouterPrefixes.TEMPLATES.value, tags=["Templates"])

TEMPLATE_HEADERS = [
    "Primary Category",
    "Detailed Category",
    "Description",
    "Date",
    "Amount",
    "Account Name",
    "Status",
    "Notes",
    "Tags",
]

_EXAMPLE_ROWS = [
    ["Food & drink", "Groceries", "Costco Wholesale", "2024-10-15", -143.27, "Checking", "Verified", "", ""],
    ["Income", "Wages", "Acme Corp Payroll", "2024-10-01", 3200.00, "Checking", "Verified", "", ""],
]

# Extra blank rows to cover with category dropdown validation beyond the example rows
_BLANK_ROWS_FOR_VALIDATION = 200


def _build_template_workbook() -> Workbook:
    """Builds an in-memory .xlsx workbook for the CSV/Excel import template."""
    primary_values = [c.value for c in PrimaryCategories]
    detailed_values = [c.value for c in DetailedCategories]

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(TEMPLATE_HEADERS)
    for row in _EXAMPLE_ROWS:
        ws.append(row)
    for col_idx in range(1, len(TEMPLATE_HEADERS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 18

    # Hidden helper sheet backing the category dropdowns
    categories_sheet = wb.create_sheet("Categories")
    categories_sheet.sheet_state = "hidden"
    categories_sheet.append(["Primary Category", "Detailed Category"])
    for idx, value in enumerate(primary_values, start=2):
        categories_sheet.cell(row=idx, column=1, value=value)
    for idx, value in enumerate(detailed_values, start=2):
        categories_sheet.cell(row=idx, column=2, value=value)

    wb.defined_names.add(DefinedName(
        "PrimaryCategoryList",
        attr_text=f"Categories!$A$2:$A${len(primary_values) + 1}",
    ))
    wb.defined_names.add(DefinedName(
        "DetailedCategoryList",
        attr_text=f"Categories!$B$2:$B${len(detailed_values) + 1}",
    ))

    primary_validation = DataValidation(type="list", formula1="=PrimaryCategoryList", allow_blank=True)
    primary_validation.error = "Select a category from the dropdown list."
    primary_validation.errorTitle = "Invalid Primary Category"
    detailed_validation = DataValidation(type="list", formula1="=DetailedCategoryList", allow_blank=True)
    detailed_validation.error = "Select a category from the dropdown list."
    detailed_validation.errorTitle = "Invalid Detailed Category"

    last_row = ws.max_row + _BLANK_ROWS_FOR_VALIDATION
    primary_validation.add(f"A2:A{last_row}")
    detailed_validation.add(f"B2:B{last_row}")
    ws.add_data_validation(primary_validation)
    ws.add_data_validation(detailed_validation)

    return wb


@router.get("/csv-import")
def download_csv_import_template() -> Response:
    """Returns a downloadable .xlsx import template with category dropdown validation."""
    wb = _build_template_workbook()
    buffer = BytesIO()
    wb.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="budgy_import_template.xlsx"'},
    )
